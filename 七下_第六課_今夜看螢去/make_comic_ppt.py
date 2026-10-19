# -*- coding: utf-8 -*-
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = r'E:\antigravity\202606_國文教材\七下_第六課_今夜看螢去\課文_今夜看螢去_情境教學簡報.pptx'
IMG_DIR = r'E:\antigravity\202606_國文教材\七下_第六課_今夜看螢去\圖片'

prs = pptx.Presentation()
prs.slide_width = 9144000  # 16:9
prs.slide_height = 5143500
blank_layout = prs.slide_layouts[6]

panels = [
    {
        'id': 1,
        'title': '第一段：荒郊尋螢',
        'text': '我很久沒有看到螢火蟲了。有一天夜晚，我專程到臺北附近人少的荒郊看螢火蟲。',
        'img': '荒郊.jpg'
    },
    {
        'id': 2,
        'title': '第二段（上）：溪水潺潺',
        'text': '我沿著山路往前走，聽到潺潺的水聲。',
        'img': '潺潺.jpg'
    },
    {
        'id': 3,
        'title': '第二段（下）：草叢閃爍',
        'text': '忽然，草叢裡飛出一隻螢火蟲，接著又飛出幾隻。牠們的光在黑暗中閃爍。',
        'img': '閃爍.jpg'
    },
    {
        'id': 4,
        'title': '第五段：舒緩飛舞',
        'text': '我喜歡看螢火蟲飛舞。牠們有時飛得舒緩，有時一下飛高、一下飛低，我會一直看著牠們。',
        'img': '舒緩.jpg'
    },
    {
        'id': 5,
        'title': '第六段：掌心相遇',
        'text': '我慢慢伸出手掌，讓螢火蟲停在掌心。牠只停一會兒，我很珍惜這短短的相遇。',
        'img': '珍惜.jpg'
    },
    {
        'id': 6,
        'title': '第七段（上）：都市遺憾',
        'text': '現在螢火蟲比以前少了。都市裡有些孩子很少看見牠們，讓我覺得遺憾。',
        'img': '遺憾.jpg'
    },
    {
        'id': 7,
        'title': '第七段（下）：善待生靈',
        'text': '我們應該善待螢火蟲，不要傷害牠們。',
        'img': '善待.jpg'
    },
    {
        'id': 8,
        'title': '第八段：美好回憶',
        'text': '這次看螢火蟲，我不只看到了閃亮的光點，也想起美好的童年回憶。',
        'img': '專程.jpg'
    }
]

def add_bg(slide, r=250, g=250, b=252):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(r, g, b)
    bg.line.color.rgb = RGBColor(r, g, b)
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)

# Slide 1: Cover
s1 = prs.slides.add_slide(blank_layout)
add_bg(s1, 245, 247, 252)

tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(8.0), Inches(3.0))
tf = tb.text_frame
p1 = tf.paragraphs[0]
p1.text = '七下 國語文 第六課 · 特教繪本情境教學'
p1.font.name = '微軟正黑體'
p1.font.size = Pt(22)
p1.font.color.rgb = RGBColor(99, 102, 241)
p1.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = '今夜看螢去'
p2.font.name = '微軟正黑體'
p2.font.size = Pt(56)
p2.font.bold = True
p2.font.color.rgb = RGBColor(30, 41, 59)
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(12)

p3 = tf.add_paragraph()
p3.text = '【米克師連環漫畫雙模式版】· 漫畫一覽 + 逐格滿版繪本'
p3.font.name = '微軟正黑體'
p3.font.size = Pt(20)
p3.font.color.rgb = RGBColor(100, 116, 139)
p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(18)

# Helper for Comic Grid Page (4 panels per page)
def make_comic_grid(slide, panel_subset, page_title):
    add_bg(slide, 248, 250, 252)
    # Header
    tb_h = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(8.8), Inches(0.6))
    tf_h = tb_h.text_frame
    p_h = tf_h.paragraphs[0]
    p_h.text = '【連環漫畫模式】' + page_title
    p_h.font.name = '微軟正黑體'
    p_h.font.size = Pt(20)
    p_h.font.bold = True
    p_h.font.color.rgb = RGBColor(79, 70, 229)
    
    # 2x2 grid positions
    coords = [
        (Inches(0.6), Inches(0.95)),
        (Inches(5.2), Inches(0.95)),
        (Inches(0.6), Inches(3.05)),
        (Inches(5.2), Inches(3.05))
    ]
    card_w = Inches(4.2)
    card_h = Inches(1.95)
    
    for i, pdata in enumerate(panel_subset):
        x, y = coords[i]
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = RGBColor(226, 232, 240)
        card.line.width = Pt(1.5)
        
        # Image
        img_p = os.path.join(IMG_DIR, pdata['img'])
        if os.path.exists(img_p):
            slide.shapes.add_picture(img_p, x + Inches(0.12), y + Inches(0.12), width=Inches(1.8), height=Inches(1.7))
        
        # Text on right of card
        tx = slide.shapes.add_textbox(x + Inches(2.0), y + Inches(0.08), Inches(2.1), Inches(1.75))
        tf_c = tx.text_frame
        tf_c.word_wrap = True
        
        # Tag
        ptag = tf_c.paragraphs[0]
        ptag.text = '第 ' + str(pdata['id']) + ' 格 · ' + pdata['title']
        ptag.font.name = '微軟正黑體'
        ptag.font.size = Pt(13)
        ptag.font.bold = True
        ptag.font.color.rgb = RGBColor(99, 102, 241)
        ptag.space_after = Pt(4)
        
        # Sentence
        ptext = tf_c.add_paragraph()
        ptext.text = pdata['text']
        ptext.font.name = '微軟正黑體'
        ptext.font.size = Pt(12)
        ptext.font.color.rgb = RGBColor(51, 65, 85)
        ptext.line_spacing = 1.15

# Slide 2: Comic Grid (Panels 1-4)
s2 = prs.slides.add_slide(blank_layout)
make_comic_grid(s2, panels[0:4], '故事起承（第 1 ~ 4 格）')

# Slide 3: Comic Grid (Panels 5-8)
s3 = prs.slides.add_slide(blank_layout)
make_comic_grid(s3, panels[4:8], '故事轉合（第 5 ~ 8 格）')

# Slide 4~11: Single Full-Picture Storybook Slides (1 panel per slide)
for pdata in panels:
    s = prs.slides.add_slide(blank_layout)
    add_bg(s, 255, 255, 255)
    
    # Header tag
    tb_tag = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(8.4), Inches(0.4))
    tf_tag = tb_tag.text_frame
    ptg = tf_tag.paragraphs[0]
    ptg.text = '【滿版大圖繪本模式】第 ' + str(pdata['id']) + ' 格 · ' + pdata['title']
    ptg.font.name = '微軟正黑體'
    ptg.font.size = Pt(16)
    ptg.font.bold = True
    ptg.font.color.rgb = RGBColor(99, 102, 241)
    
    # Picture (Centered in top area)
    img_p = os.path.join(IMG_DIR, pdata['img'])
    if os.path.exists(img_p):
        s.shapes.add_picture(img_p, Inches(2.2), Inches(0.75), width=Inches(5.6))
    
    # Text Box below picture
    tb_txt = s.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(8.4), Inches(1.2))
    tf_t = tb_txt.text_frame
    tf_t.word_wrap = True
    pt = tf_t.paragraphs[0]
    pt.text = pdata['text']
    pt.font.name = '微軟正黑體'
    pt.font.size = Pt(22)
    pt.font.bold = True
    pt.font.color.rgb = RGBColor(30, 41, 59)
    pt.alignment = PP_ALIGN.CENTER
    pt.line_spacing = 1.25

# Slide 12: Summary
s12 = prs.slides.add_slide(blank_layout)
add_bg(s12, 248, 250, 252)
tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.4), Inches(4.2))
tf12 = tb12.text_frame
tf12.word_wrap = True

pt12 = tf12.paragraphs[0]
pt12.text = '【今夜看螢去】全篇漫畫情境與學習總結'
pt12.font.name = '微軟正黑體'
pt12.font.size = Pt(26)
pt12.font.bold = True
pt12.font.color.rgb = RGBColor(40, 70, 120)
pt12.space_after = Pt(20)

p_sum1 = tf12.add_paragraph()
p_sum1.text = '故事主軸回顧：'
p_sum1.font.name = '微軟正黑體'
p_sum1.font.size = Pt(20)
p_sum1.font.bold = True
p_sum1.font.color.rgb = RGBColor(79, 70, 229)

p_sum1_desc = tf12.add_paragraph()
p_sum1_desc.text = '專程荒郊尋螢 → 山路潺潺溪水 → 草叢閃爍微光 → 舒緩夜空飛舞 → 掌心短暫相遇珍惜 → 城市螢光少而遺憾 → 溫柔善待小生靈 → 銘記美好回憶'
p_sum1_desc.font.name = '微軟正黑體'
p_sum1_desc.font.size = Pt(17)
p_sum1_desc.font.color.rgb = RGBColor(51, 65, 85)
p_sum1_desc.space_after = Pt(20)
p_sum1_desc.line_spacing = 1.3

p_sum2 = tf12.add_paragraph()
p_sum2.text = '核心教學理念：'
p_sum2.font.name = '微軟正黑體'
p_sum2.font.size = Pt(20)
p_sum2.font.bold = True
p_sum2.font.color.rgb = RGBColor(16, 185, 129)

p_sum2_desc = tf12.add_paragraph()
p_sum2_desc.text = '透過連貫的手繪色鉛筆分鏡，幫助特教學童具象化理解課文中的大自然觀察與愛護生命的溫柔情感。'
p_sum2_desc.font.name = '微軟正黑體'
p_sum2_desc.font.size = Pt(17)
p_sum2_desc.font.color.rgb = RGBColor(51, 65, 85)

prs.save(PPTX_PATH)
print('PPTX generated successfully!')
