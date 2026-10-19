# -*- coding: utf-8 -*-
import os, sys
from PIL import Image
import pptx
import numpy as np

BASE_DIR = r'E:ntigravity606_國文教材\七下_第六課_今夜看螢去'
CARDS_DIR = os.path.join(BASE_DIR, r'圖片\難語詞')
IMG_DIR = os.path.join(BASE_DIR, r'圖片')
PPTX_PATH = os.path.join(BASE_DIR, r'語詞_今夜看螢去_難語詞簡報.pptx')

WORDS = ['專程', '荒郊', '潺潺', '閃爍', '舒緩', '珍惜', '善待', '遺憾']
TARGET_WORDS = ['舒緩', '珍惜', '善待', '遺憾']

def find_comic_box(im):
    arr = np.array(im.convert('RGB'))
    # Search in middle region (Column 2 comic)
    sub = arr[650:1250, 800:1650]
    gray = np.mean(sub, axis=2)
    mask = gray < 70
    rows = np.sum(mask, axis=1)
    cols = np.sum(mask, axis=0)
    
    r_idx = np.where(rows > 180)[0]
    c_idx = np.where(cols > 120)[0]
    
    if len(r_idx) >= 2 and len(c_idx) >= 2:
        min_y = r_idx[0] + 650
        max_y = r_idx[-1] + 650
        min_x = c_idx[0] + 800
        max_x = c_idx[-1] + 800
        w = max_x - min_x
        h = max_y - min_y
        if 400 <= w <= 800 and 250 <= h <= 600:
            return (min_x - 5, min_y - 5, max_x + 5, max_y + 5)
    # Default fallback box based on 荒郊.png
    return (915, 740, 1505, 1160)

def process():
    print('Checking cards in:', CARDS_DIR)
    for w in TARGET_WORDS:
        p = os.path.join(CARDS_DIR, f'{w}.png')
        if not os.path.exists(p):
            print(f'[Skip] {w}.png not found.')
            continue
        im = Image.open(p)
        print(f'{w}.png original size: {im.size}')
        if im.size != (2666, 1500):
            im = im.resize((2666, 1500), Image.Resampling.LANCZOS)
            im.save(p)
            print(f'-> Resized {w}.png to (2666, 1500)')
        
        crop_box = find_comic_box(im)
        cropped = im.crop(crop_box)
        if cropped.mode in ('RGBA', 'P'):
            cropped = cropped.convert('RGB')
        out_jpg = os.path.join(IMG_DIR, f'{w}.jpg')
        cropped.save(out_jpg, 'JPEG', quality=95)
        print(f'-> Cropped {w}.jpg: bbox={crop_box}, size={cropped.size}, saved to {out_jpg}')

    print('Updating PPTX...')
    prs = pptx.Presentation(PPTX_PATH)
    for i, w in enumerate(WORDS):
        card_p = os.path.join(CARDS_DIR, f'{w}.png')
        if os.path.exists(card_p) and i < len(prs.slides):
            slide = prs.slides[i]
            sp_to_remove = [s for s in slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE]
            for sp in sp_to_remove:
                sp._element.getparent().remove(sp._element)
            slide.shapes.add_picture(card_p, 0, 0, prs.slide_width, prs.slide_height)
            print(f'-> Slide {i+1} ({w}) updated in PPTX.')
    
    prs.save(PPTX_PATH)
    print('PPTX saved successfully!')

if __name__ == '__main__':
    process()
